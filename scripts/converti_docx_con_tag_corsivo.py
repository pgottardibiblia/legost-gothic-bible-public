#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Created on Thu Jul  4 08:36:20 2024

@author: pierandrea
"""

#Script per convertire file .docx in file .txt con tag per il corsivo
def converti_docx_con_tag_corsivo(input_path, output_path):
    """
    Legge un file .docx, aggiunge i tag <hi rend="italic"> attorno alle parti 
    in corsivo e salva il risultato come file .txt
    
    Args:
        input_path (str): Percorso del file .docx di input
        output_path (str): Percorso del file .txt di output
    """
    import sys
    from docx import Document
    
    try:
        # Carica il documento
        print(f"Caricamento del file: {input_path}")
        doc = Document(input_path)
        
        # Apri il file di output
        with open(output_path, 'w', encoding='utf-8') as f:
            
            # Processa tutti i paragrafi del documento principale
            for paragrafo in doc.paragraphs:
                risultato = ""
                
                # Analizza ogni "run" (porzione di testo con formattazione uniforme)
                for run in paragrafo.runs:
                    testo = run.text
                    
                    # Se il run è in corsivo, aggiungi i tag
                    if run.italic:
                        risultato += f'<hi rend="italic">{testo}</hi>'
                    else:
                        risultato += testo
                
                # Scrivi solo i paragrafi non vuoti
                if risultato.strip():
                    f.write(risultato + '\n')
            
            # Processa anche il testo nelle tabelle (se presenti)
            for tabella in doc.tables:
                for riga in tabella.rows:
                    for cella in riga.cells:
                        for paragrafo in cella.paragraphs:
                            risultato = ""
                            
                            for run in paragrafo.runs:
                                testo = run.text
                                if run.italic:
                                    risultato += f'<hi rend="italic">{testo}</hi>'
                                else:
                                    risultato += testo
                            
                            if risultato.strip():
                                f.write(risultato + '\n')
        
        print(f"✓ File elaborato con successo!")
        print(f"✓ Output salvato in: {output_path}")
        
    except FileNotFoundError:
        print(f"✗ Errore: Il file '{input_path}' non è stato trovato.")
        sys.exit(1)
    except Exception as e:
        print(f"✗ Errore durante l'elaborazione: {e}")
        sys.exit(1)
